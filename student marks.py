import os
import tkinter as tk
from tkinter import messagebox, ttk
import openpyxl

EXCEL_FILE = "student_results.xlsx"
HEADERS = [
    "Name",
    "Roll No.",
    "Class",
    "Subject 1",
    "Subject 2",
    "Subject 3",
    "Subject 4",
    "Subject 5",
    "Total Marks",
    "Percentage",
    "Result",
]


def initialize_excel():
    if not os.path.exists(EXCEL_FILE):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Results"
        ws.append(HEADERS)
        wb.save(EXCEL_FILE)


class StudentManagementApp(tk.Tk):
    
    def __init__(self):
        super().__init__()
        self.title("Student Result Management System")
        self.geometry("800x600")
        self.configure(bg="#f4f6f9")

        initialize_excel()

        title_frame = tk.Frame(self, bg="#2c3e50", pady=15)
        title_frame.pack(fill=tk.X)
        title_label = tk.Label(
            title_frame,
            text="STUDENT RESULT MANAGEMENT SYSTEM",
            font=("Helvetica", 18, "bold"),
            fg="white",
            bg="#2c3e50",
        )
        title_label.pack()

        nav_frame = tk.Frame(self, bg="#34495e")
        nav_frame.pack(fill=tk.X)

        btn_style = {
            "font": ("Helvetica", 11, "bold"),
            "fg": "white",
            "bg": "#34495e",
            "activebackground": "#1abc9c",
            "activeforeground": "white",
            "bd": 0,
            "padx": 20,
            "pady": 10,
        }

        tk.Button(
            nav_frame,
            text="Add Student",
            command=self.show_add_student,
            **btn_style,
        ).pack(side=tk.LEFT, expand=True, fill=tk.X)
        tk.Button(
            nav_frame,
            text="Get Result",
            command=self.show_get_result,
            **btn_style,
        ).pack(side=tk.LEFT, expand=True, fill=tk.X)
        tk.Button(
            nav_frame,
            text="Show All Results",
            command=self.show_all_results,
            **btn_style,
        ).pack(side=tk.LEFT, expand=True, fill=tk.X)

        self.container = tk.Frame(self, bg="#f4f6f9", padx=20, pady=20)
        self.container.pack(fill=tk.BOTH, expand=True)

        self.show_add_student()

    def clear_container(self):
        for widget in self.container.winfo_children():
            widget.destroy()

    def show_add_student(self):
        self.clear_container()

        lbl_title = tk.Label(
            self.container,
            text="Add Student Details & Marks",
            font=("Helvetica", 14, "bold"),
            bg="#f4f6f9",
            fg="#2c3e50",
        )
        lbl_title.pack(anchor=tk.W, pady=(0, 15))

        form_frame = tk.Frame(self.container, bg="#f4f6f9")
        form_frame.pack(fill=tk.X)

        labels = [
            "Name:",
            "Roll No.:",
            "Class:",
            "Subject 1 Marks:",
            "Subject 2 Marks:",
            "Subject 3 Marks:",
            "Subject 4 Marks:",
            "Subject 5 Marks:",
        ]
        self.entries = {}

        for i, text in enumerate(labels):
            lbl = tk.Label(
                form_frame,
                text=text,
                font=("Helvetica", 11),
                bg="#f4f6f9",
                anchor="w",
            )
            lbl.grid(row=i, column=0, sticky=tk.W, pady=5, padx=5)

            ent = tk.Entry(form_frame, font=("Helvetica", 11), width=30)
            ent.grid(row=i, column=1, pady=5, padx=5)
            self.entries[text] = ent

        btn_save = tk.Button(
            self.container,
            text="Save Record",
            font=("Helvetica", 12, "bold"),
            bg="#27ae60",
            fg="white",
            padx=15,
            pady=8,
            bd=0,
            command=self.save_student,
        )
        btn_save.pack(anchor=tk.W, pady=20)

    def save_student(self):
        try:
            name = self.entries["Name:"].get().strip()
            roll_no = self.entries["Roll No.:"].get().strip()
            student_cls = self.entries["Class:"].get().strip()

            if not name or not roll_no or not student_cls:
                messagebox.showerror(
                    "Error", "Name, Roll No., and Class are required fields."
                )
                return

            marks = []
            for i in range(1, 6):
                m = float(self.entries[f"Subject {i} Marks:"].get().strip())
                if not (0 <= m <= 100):
                    raise ValueError(f"Subject {i} marks must be between 0 and 100.")
                marks.append(m)

            total = sum(marks)
            percentage = total / 5.0
            result = (
                "Pass"
                if percentage >= 40 and all(m >= 33 for m in marks)
                else "Fail"
            )

            wb = openpyxl.load_workbook(EXCEL_FILE)
            ws = wb.active
            for row in ws.iter_rows(min_row=2, values_only=True):
                if str(row[1]) == str(roll_no):
                    messagebox.showerror(
                        "Error",
                        f"Student with Roll No. {roll_no} already exists!",
                    )
                    wb.close()
                    return

            record = [
                name,
                roll_no,
                student_cls,
                *marks,
                total,
                f"{percentage:.1f}%",
                result,
            ]
            ws.append(record)
            wb.save(EXCEL_FILE)
            wb.close()

            messagebox.showinfo("Success", "Student record saved successfully!")
            self.show_add_student()  # Reset form

        except ValueError as ve:
            messagebox.showerror(
                "Input Error", f"Invalid input: {ve}\nPlease enter valid numeric marks."
            )

    # --- 2. GET RESULT VIEW ---
    def show_get_result(self):
        self.clear_container()

        lbl_title = tk.Label(
            self.container,
            text="Search Student Result",
            font=("Helvetica", 14, "bold"),
            bg="#f4f6f9",
            fg="#2c3e50",
        )
        lbl_title.pack(anchor=tk.W, pady=(0, 15))

        search_frame = tk.Frame(self.container, bg="#f4f6f9")
        search_frame.pack(fill=tk.X, pady=10)

        tk.Label(
            search_frame,
            text="Enter Roll No.:",
            font=("Helvetica", 11),
            bg="#f4f6f9",
        ).pack(side=tk.LEFT, padx=(0, 10))

        self.search_entry = tk.Entry(
            search_frame, font=("Helvetica", 11), width=15
        )
        self.search_entry.pack(side=tk.LEFT, padx=5)

        btn_search = tk.Button(
            search_frame,
            text=" Get Result",
            font=("Helvetica", 10, "bold"),
            bg="#2980b9",
            fg="white",
            padx=10,
            pady=3,
            bd=0,
            command=self.search_result,
        )
        btn_search.pack(side=tk.LEFT, padx=10)

        self.result_display_frame = tk.Frame(self.container, bg="#f4f6f9")
        self.result_display_frame.pack(fill=tk.BOTH, expand=True, pady=20)

    def search_result(self):
        for widget in self.result_display_frame.winfo_children():
            widget.destroy()

        roll_no = self.search_entry.get().strip()
        if not roll_no:
            messagebox.showwarning("Warning", "Please enter a Roll No.")
            return

        wb = openpyxl.load_workbook(EXCEL_FILE)
        ws = wb.active
        matched_student = None

        for row in ws.iter_rows(min_row=2, values_only=True):
            if str(row[1]) == roll_no:
                matched_student = row
                break
        wb.close()

        if matched_student:
            cols = ["Name", "Roll No.", "Class", "Total Marks", "Percentage", "Result"]
            display_values = [
                matched_student[0],
                matched_student[1],
                matched_student[2],
                matched_student[8],
                matched_student[9],
                matched_student[10],
            ]

            card = tk.Frame(
                self.result_display_frame,
                bg="white",
                bd=1,
                relief=tk.SOLID,
                padx=20,
                pady=20,
            )
            card.pack(anchor=tk.NW, fill=tk.X)

            tk.Label(
                card,
                text="Student Record Found",
                font=("Helvetica", 12, "bold"),
                bg="white",
                fg="#27ae60",
            ).pack(anchor=tk.W, pady=(0, 10))

            for col, val in zip(cols, display_values):
                row_f = tk.Frame(card, bg="white")
                row_f.pack(fill=tk.X, pady=3)
                tk.Label(
                    row_f,
                    text=f"{col}:",
                    font=("Helvetica", 11, "bold"),
                    bg="white",
                    width=15,
                    anchor="w",
                ).pack(side=tk.LEFT)
                tk.Label(
                    row_f,
                    text=str(val),
                    font=("Helvetica", 11),
                    bg="white",
                    anchor="w",
                ).pack(side=tk.LEFT)
        else:
            msg_label = tk.Label(
                self.result_display_frame,
                text=" Student record not found.",
                font=("Helvetica", 12, "bold"),
                fg="#e74c3c",
                bg="#f4f6f9",
            )
            msg_label.pack(anchor=tk.W)

    # --- 3. SHOW ALL RESULTS VIEW ---
    def show_all_results(self):
        self.clear_container()

        lbl_title = tk.Label(
            self.container,
            text="All Student Results",
            font=("Helvetica", 14, "bold"),
            bg="#f4f6f9",
            fg="#2c3e50",
        )
        lbl_title.pack(anchor=tk.W, pady=(0, 15))

        # Setup Treeview
        tree_columns = (
            "Name",
            "Roll No.",
            "Class",
            "Total Marks",
            "Percentage",
            "Result",
        )
        tree = ttk.Treeview(
            self.container, columns=tree_columns, show="headings"
        )

        for col in tree_columns:
            tree.heading(col, text=col)
            tree.column(col, anchor=tk.CENTER, width=110)

        # Scrollbar
        scrollbar = ttk.Scrollbar(
            self.container, orient=tk.VERTICAL, command=tree.yview
        )
        tree.configure(yscroll=scrollbar.set)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        tree.pack(fill=tk.BOTH, expand=True)

        # Populate records
        wb = openpyxl.load_workbook(EXCEL_FILE)
        ws = wb.active

        for row in ws.iter_rows(min_row=2, values_only=True):
            display_data = (row[0], row[1], row[2], row[8], row[9], row[10])
            tree.insert("", tk.END, values=display_data)

        wb.close()


if __name__ == "__main__":
    app = StudentManagementApp()
    app.mainloop()
